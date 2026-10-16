"""csv_a_excel.py — Kit eQ, skill "reportes".

Convierte el CSV que deja consultar.ps1 en un Excel prolijo y verificable:
  - Hoja "Datos": encabezado en negrita, filtros, columnas anchas, fila de TOTALES
    con fórmulas =SUBTOTAL(9; ...) (los totales los cierra Excel, no Claude).
  - Hoja "Control": de dónde salió (vista y consulta), cuándo, cuántas filas,
    y los totales calculados acá para cruzarlos contra los de Excel.

Uso:
  py csv_a_excel.py --csv resultados/x.csv --xlsx resultados/x.xlsx \
     --origen "dbo.MiVista" --consulta "SELECT ..." [--sumar Col1,Col2]

No pisa archivos: si el .xlsx ya existe, frena.
Necesita: Python 3 y openpyxl (los instala la skill "configurar").
"""
import argparse
import csv
import datetime as dt
import os
import re
import sys
from decimal import Decimal, InvalidOperation

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("FALTA openpyxl: corré la skill 'configurar' (o: py -m pip install --user openpyxl).")
    sys.exit(3)


# Hasta 15 dígitos enteros: Excel no guarda más sin perder precisión.
NUMERO = re.compile(r"-?\d{1,15}(\.\d+)?")


def a_numero(texto):
    """Convierte '1234.5' (punto decimal, como lo deja consultar.ps1) a número.
    Devuelve None si no es número. Los códigos con cero adelante ('00123') quedan como texto."""
    if texto is None:
        return None
    t = texto.strip()
    if not t or "," in t:
        return None
    if not NUMERO.fullmatch(t):
        return None  # CBU, CUIT largos, códigos con letras o con E: quedan como texto
    sin_signo = t.lstrip("-")
    if len(sin_signo) > 1 and sin_signo.startswith("0") and not sin_signo.startswith("0."):
        return None
    try:
        d = Decimal(t)
    except InvalidOperation:
        return None
    if not d.is_finite():
        return None
    return int(d) if "." not in t else float(d)


def a_fecha(texto):
    """'2026-10-07' o '2026-10-07 00:00:00' -> fecha. None si no es fecha."""
    t = (texto or "").strip()
    for formato in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(t, formato)
        except ValueError:
            continue
    return None


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", required=True)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--origen", default="(sin dato)")
    ap.add_argument("--consulta", default="(sin dato)")
    ap.add_argument("--sumar", default="", help="Columnas a totalizar, separadas por coma. Vacío = todas las numéricas.")
    a = ap.parse_args()

    if os.path.exists(a.xlsx):
        print(f"FRENADO: '{a.xlsx}' ya existe. No piso archivos: elegí otro nombre (por ejemplo _v02).")
        sys.exit(7)

    with open(a.csv, encoding="utf-8-sig", newline="") as f:
        filas = list(csv.reader(f, delimiter=";"))
    if not filas:
        print("FRENADO: el CSV está vacío (ni siquiera tiene encabezado).")
        sys.exit(6)
    encabezado, datos = filas[0], filas[1:]

    # Detectar columnas numéricas (todas sus celdas no vacías son número)
    numericas = []
    for j, _ in enumerate(encabezado):
        valores = [r[j] for r in datos if j < len(r) and r[j].strip()]
        if valores and all(a_numero(v) is not None for v in valores):
            numericas.append(j)
    fechas = []
    for j, _ in enumerate(encabezado):
        valores = [r[j] for r in datos if j < len(r) and r[j].strip()]
        if j not in numericas and valores and all(a_fecha(v) is not None for v in valores):
            fechas.append(j)
    pedidas = [s.strip() for s in a.sumar.split(",") if s.strip()]
    if pedidas:
        faltan = [p for p in pedidas if p not in encabezado]
        if faltan:
            print(f"FRENADO: no existen las columnas a sumar: {', '.join(faltan)}")
            sys.exit(6)
        a_sumar = [encabezado.index(p) for p in pedidas]
    else:
        a_sumar = numericas

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(encabezado)
    for r in datos:
        ws.append([
            a_numero(v) if j in numericas else (a_fecha(v) if j in fechas and v.strip() else v)
            for j, v in enumerate(r)
        ])

    negrita = Font(bold=True)
    for celda in ws[1]:
        celda.font = negrita
        celda.fill = PatternFill("solid", fgColor="DDEBF7")
    ultima = len(datos) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezado))}{ultima}"
    ws.freeze_panes = "A2"

    totales_py = {}
    if datos and a_sumar:
        fila_tot = ultima + 2
        ws.cell(row=fila_tot, column=1, value="TOTAL (fórmula)").font = negrita
        for j in a_sumar:
            col = get_column_letter(j + 1)
            c = ws.cell(row=fila_tot, column=j + 1, value=f"=SUBTOTAL(9,{col}2:{col}{ultima})")
            c.font = negrita
            totales_py[encabezado[j]] = float(sum(
                (Decimal(r[j].strip()) for r in datos if j < len(r) and a_numero(r[j]) is not None),
                Decimal(0),
            ))

    for j, nombre in enumerate(encabezado, start=1):
        largo = max([len(str(nombre))] + [len(str(r[j - 1])) for r in datos[:500] if j - 1 < len(r)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, largo + 2), 50)
    for j in numericas:
        # Importes (lo que se suma o lo que tiene decimales): 1.234,56. Números enteros
        # que no se suman (nro. de pedido, código): sin decimales, para que no salga "1,00".
        con_decimales = any("." in r[j] for r in datos if j < len(r))
        formato = "#,##0.00" if (j in a_sumar or con_decimales) else "0"
        for fila in ws.iter_rows(min_row=2, min_col=j + 1, max_col=j + 1):
            for c in fila:
                c.number_format = formato
    for j in fechas:
        for fila in ws.iter_rows(min_row=2, max_row=ultima, min_col=j + 1, max_col=j + 1):
            for c in fila:
                c.number_format = "dd/mm/yyyy"

    ctl = wb.create_sheet("Control")
    ctl.append(["Dato", "Valor"])
    ctl.append(["Origen (vista)", a.origen])
    ctl.append(["Consulta", a.consulta])
    ctl.append(["Generado", dt.datetime.now().strftime("%d/%m/%Y %H:%M")])
    ctl.append(["Filas de datos", len(datos)])
    ctl.append(["Columnas", len(encabezado)])
    for nombre, total in totales_py.items():
        ctl.append([f"Total calculado de {nombre} (cruzar con la fila TOTAL de Datos)", total])
    for celda in ctl[1]:
        celda.font = negrita
    ctl.column_dimensions["A"].width = 60
    ctl.column_dimensions["B"].width = 60

    os.makedirs(os.path.dirname(os.path.abspath(a.xlsx)), exist_ok=True)
    wb.save(a.xlsx)
    print("OK")
    print(f"Filas: {len(datos)}")
    print(f"Columnas: {len(encabezado)}")
    if not datos:
        print("AVISO SIN FILAS: la consulta no trajo datos. Revisá el filtro (período, cliente) antes de entregar.")
    else:
        print("Totales con fórmula en: " + (", ".join(encabezado[j] for j in a_sumar) if a_sumar else "(ninguna columna numérica)"))
    for nombre, total in totales_py.items():
        print(f"Control {nombre}: {total}")
    print(f"Archivo: {a.xlsx}")


if __name__ == "__main__":
    main()
